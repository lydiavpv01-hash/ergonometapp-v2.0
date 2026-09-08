from io import BytesIO
import base64

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

from services.matrix_excel import export_nom_excel as _base_export


NAVY = '153D63'
STRONG_BLUE = '215E99'
SKY = 'A5C9EB'
WHITE = 'FFFFFF'
BLACK = '000000'
GREEN = '00FF00'
YELLOW = 'FFFF00'
RED = 'FF0000'
MAGENTA = 'FF00FF'

NORM_ACTIONS = {
    'bajo': 'Sólo se requiere dar seguimiento a los grupos más vulnerables, como mujeres en periodo de gestación o trabajadores menores de edad.',
    'medio': 'Se debe examinar las tareas con mayor detalle, mediante la aplicación de una evaluación específica, o bien implantar medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.',
    'alto': 'Se requiere una acción rápida, por lo que se deben establecer medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.',
    'muy alto': 'Se deben detener las actividades e implementar medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.',
}


def _field_value(payload, *needles):
    needles = [str(n).lower() for n in needles]
    for f in (payload.get('snapshot') or {}).get('fields', []) or []:
        corpus = ' '.join(str(f.get(k, '')) for k in ('name', 'id', 'label')).lower()
        if any(n in corpus for n in needles):
            v = f.get('value')
            if v not in (None, ''):
                return v
    return ''


def _upload(payload, field):
    for u in payload.get('uploads') or []:
        if (u.get('field') or '') == field and u.get('data'):
            return u
    return None


def _bytes_from_upload(upload):
    if not upload:
        return None
    data = upload.get('data') or ''
    if ',' in data:
        data = data.split(',', 1)[1]
    try:
        return base64.b64decode(data)
    except Exception:
        return None


def _add_logo(ws, upload, anchor, max_w=132, max_h=50):
    raw = _bytes_from_upload(upload)
    if not raw:
        return
    bio = BytesIO(raw)
    try:
        with PILImage.open(bio) as pil:
            w, h = pil.size
        bio.seek(0)
        img = XLImage(bio)
        scale = min(max_w / max(w, 1), max_h / max(h, 1), 1.0)
        img.width = max(1, w * scale)
        img.height = max(1, h * scale)
        ws.add_image(img, anchor)
    except Exception:
        return


def _navy_border(existing=None):
    def side(style='thin'):
        return Side(style=style, color=NAVY)
    if existing:
        return Border(
            left=side(existing.left.style or 'thin') if existing.left and existing.left.style else side(),
            right=side(existing.right.style or 'thin') if existing.right and existing.right.style else side(),
            top=side(existing.top.style or 'thin') if existing.top and existing.top.style else side(),
            bottom=side(existing.bottom.style or 'thin') if existing.bottom and existing.bottom.style else side(),
        )
    return Border(left=side(), right=side(), top=side(), bottom=side())


def _apply_border_palette(ws):
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for c in row:
            b = c.border
            if any((getattr(b, side).style if getattr(b, side, None) else None) for side in ('left','right','top','bottom')):
                c.border = _navy_border(b)


def _style_merged_range(ws, rng, fill, font_color=BLACK, bold=True, size=9):
    top = ws[rng.split(':')[0]]
    top.fill = PatternFill('solid', fgColor=fill)
    top.font = Font(name='Arial', size=size, bold=bold, color=font_color)
    top.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws[rng]:
        for c in row:
            c.border = _navy_border(c.border)


def _find_merged(ws, coord):
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return str(rng)
    return None


def _style_sections(ws):
    major_tokens = (
        'EVALUACIÓN DE RIESGO POR MANEJO MANUAL',
        'EVALUACIÓN DE RIESGO POR EMPUJE',
        'VALORES PARA LA DETERMINACIÓN DEL GRADO DE RIESGO',
        'FUNDAMENTO DE LA PUNTUACIÓN Y EVIDENCIA',
    )
    sub_tokens = (
        'Información General', 'Condiciones de la actividad', 'Movimientos que realiza',
        'FACTORES DE RIESGO', 'LEVANTAR', 'TRANSPORTAR', 'EQUIPO',
        'EQUIPO PEQUEÑO', 'EQUIPO MEDIANO', 'EQUIPO GRANDE',
        'RODANDO', 'GIRANDO SOBRE SU BASE', 'ARRASTRANDO/JALANDO O DESLIZANDO',
        'NIVEL DE RIESGO', 'PUNTAJE TOTAL', 'ACCIONES', 'COLOR', 'VALOR',
        'FACTOR / ACTIVIDAD', 'CONDICIÓN SELECCIONADA', 'JUSTIFICACIÓN DEL EVALUADOR', 'EVIDENCIA'
    )
    seen = set()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for c in row:
            text = str(c.value or '').strip()
            if not text:
                continue
            rng = _find_merged(ws, c.coordinate) or c.coordinate
            if rng in seen:
                continue
            if any(t in text for t in major_tokens):
                _style_merged_range(ws, rng, STRONG_BLUE, WHITE, True, 9)
                seen.add(rng)
            elif text in sub_tokens:
                _style_merged_range(ws, rng, SKY, NAVY, True, 8)
                seen.add(rng)


def _set_general_information(ws, payload):
    ws['A3'] = 'Razón Social:'
    ws['B3'] = _field_value(payload, 'razón social', 'razon social', 'empresa')
    ws['A4'] = 'Área:'
    ws['B4'] = _field_value(payload, 'area', 'área')
    ws['A5'] = 'Subárea:'
    ws['B5'] = _field_value(payload, 'subarea', 'subárea')
    ws['A8'] = 'Actividad/tarea:'
    ws['B8'] = _field_value(payload, 'actividad')
    ws['A9'] = 'Descripción:'
    ws['B9'] = _field_value(payload, 'descripcion', 'descripción')
    for cell in ('A3','A4','A5','A8','A9'):
        ws[cell].font = Font(name='Arial', size=9, bold=True, color=NAVY)
    for cell in ('B3','B4','B5','B8','B9'):
        ws[cell].font = Font(name='Arial', size=9, color=BLACK)
        ws[cell].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def _title_and_logos(ws, payload, filename):
    old_title = ws['A1'].value or ''
    try:
        ws.unmerge_cells('A1:AF1')
    except Exception:
        pass
    for rng in ('A1:F1','G1:Z1','AA1:AF1'):
        try:
            ws.merge_cells(rng)
        except Exception:
            pass
    ws.row_dimensions[1].height = 54
    ws['A1'] = ''
    ws['G1'] = old_title
    ws['AA1'] = ''
    _style_merged_range(ws, 'G1:Z1', NAVY, WHITE, True, 11)
    _style_merged_range(ws, 'A1:F1', WHITE, BLACK, False, 8)
    _style_merged_range(ws, 'AA1:AF1', WHITE, BLACK, False, 8)
    _add_logo(ws, _upload(payload, 'logo_cliente'), 'A1')
    _add_logo(ws, _upload(payload, 'logo_rfranyutti'), 'AA1')
    code = 'Fo-NSTPS-61' if '61_' in filename else ('Fo-NSTPS-63' if '63_' in filename else 'Fo-NSTPS-62')
    try:
        ws.merge_cells('AA2:AF2')
    except Exception:
        pass
    ws['AA2'] = code
    _style_merged_range(ws, 'AA2:AF2', STRONG_BLUE, WHITE, True, 8)


def _normative_actions(ws):
    for r in range(1, ws.max_row + 1):
        value = str(ws[f'A{r}'].value or '').lower().replace('–','-').strip()
        key = None
        if value.startswith('muy alto'):
            key = 'muy alto'
        elif value.startswith('alto'):
            key = 'alto'
        elif value.startswith('medio'):
            key = 'medio'
        elif value.startswith('bajo'):
            key = 'bajo'
        if not key:
            continue
        if ws[f'K{r}'].value is not None:
            ws[f'K{r}'] = NORM_ACTIONS[key]
            ws[f'K{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 38)


def _risk_colors(ws):
    colors = {'bajo': (GREEN, BLACK), 'medio': (YELLOW, BLACK), 'alto': (RED, WHITE), 'muy alto': (MAGENTA, WHITE)}
    for r in range(1, ws.max_row + 1):
        label = str(ws[f'A{r}'].value or '').lower().replace('–','-').strip()
        key = None
        if label.startswith('muy alto'): key = 'muy alto'
        elif label.startswith('alto'): key = 'alto'
        elif label.startswith('medio'): key = 'medio'
        elif label.startswith('bajo'): key = 'bajo'
        if not key:
            continue
        fill, font_color = colors[key]
        if ws[f'K{r}'].value is None:
            continue
        for c in ws[r]:
            if c.column <= 32:
                c.fill = PatternFill('solid', fgColor=fill)
                c.font = Font(name='Arial', size=8, bold=(c.column <= 4), color=font_color)
                c.border = _navy_border(c.border)
                c.alignment = Alignment(horizontal='left' if c.column >= 11 else 'center', vertical='center', wrap_text=True)


def _footer(ws):
    left = 'Documento original'
    right = 'Página &P I &N'
    ws.oddFooter.left.text = left
    ws.oddFooter.right.text = right
    ws.evenFooter.left.text = left
    ws.evenFooter.right.text = right
    ws.oddFooter.left.size = 9
    ws.oddFooter.right.size = 9
    ws.evenFooter.left.size = 9
    ws.evenFooter.right.size = 9


def export_nom_excel(matrix, payload):
    stream, filename = _base_export(matrix, payload)
    wb = load_workbook(stream)
    ws = wb.active
    _title_and_logos(ws, payload, filename)
    _set_general_information(ws, payload)
    _style_sections(ws)
    _normative_actions(ws)
    _risk_colors(ws)
    _apply_border_palette(ws)
    _footer(ws)
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out, filename
