from io import BytesIO
import base64
from copy import copy

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.drawing.image import Image as XLImage


BLACK = '000000'
WHITE = 'FFFFFF'
GRAY = 'B7B7B7'
LIGHT_GRAY = 'E7E6E6'
GREEN = '00FF00'
YELLOW = 'FFFF00'
RED = 'FF0000'
MAGENTA = 'D000D0'

THIN = Side(style='thin', color=BLACK)
MEDIUM = Side(style='medium', color=BLACK)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RISK_STYLES = {
    'Verde': (GREEN, BLACK),
    'Amarillo': (YELLOW, BLACK),
    'Naranja': (YELLOW, BLACK),
    'Rojo': (RED, WHITE),
    'Magenta': (MAGENTA, WHITE),
    'Morado': (MAGENTA, WHITE),
}


def _merge(ws, rng, value=None, fill=None, font=None, alignment=None, border=BORDER):
    ws.merge_cells(rng)
    cell = ws[rng.split(':')[0]]
    if value is not None:
        cell.value = value
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    if border:
        # Apply borders to all cells in merged range so Excel/print retains the box.
        for row in ws[rng]:
            for c in row:
                c.border = border
    return cell


def _set(ws, cell, value=None, fill=None, font=None, alignment=None, border=BORDER):
    c = ws[cell]
    c.value = value
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    if font:
        c.font = font
    if alignment:
        c.alignment = alignment
    if border:
        c.border = border
    return c


def _risk_fill(ws, cell, color_name, value=None):
    fill, font_color = RISK_STYLES.get(color_name or '', (WHITE, BLACK))
    c = ws[cell]
    c.value = color_name or '—' if value is None else value
    c.fill = PatternFill('solid', fgColor=fill)
    c.font = Font(name='Arial', size=9, bold=True, color=font_color)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = BORDER


def _risk_label(score):
    try:
        score = int(score or 0)
    except Exception:
        score = 0
    if score >= 21:
        return 'MUY ALTO - INACEPTABLE', 'Se deben detener las actividades e implementar medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.'
    if score >= 13:
        return 'ALTO - SIGNIFICATIVO', 'Se requiere una acción rápida, por lo que se deben establecer medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.'
    if score >= 5:
        return 'MEDIO - POSIBLE', 'Se debe examinar la tarea con mayor detalle e implantar medidas de control a corto plazo.'
    return 'BAJO - ACEPTABLE', 'No se requieren acciones correctivas; mantener seguimiento y control.'


def _base_sheet(title):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.sheet_view.showGridLines = False
    # Approximate the proportions of the client matrices supplied by the user.
    widths = {
        'A': 14, 'B': 12, 'C': 12, 'D': 12, 'E': 12, 'F': 12, 'G': 12,
        'H': 12, 'I': 12, 'J': 12, 'K': 12, 'L': 12, 'M': 12,
        'N': 12, 'O': 12, 'P': 12, 'Q': 12, 'R': 10, 'S': 10, 'T': 10,
        'U': 10, 'V': 10, 'W': 12, 'X': 12, 'Y': 12, 'Z': 10,
        'AA': 10, 'AB': 10, 'AC': 12, 'AD': 12, 'AE': 12, 'AF': 10,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for r in range(1, 80):
        ws.row_dimensions[r].height = 19
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    return wb, ws


def _general_header(ws, title, meta, payload=None):
    title_font = Font(name='Arial', size=12, bold=True)
    section_font = Font(name='Arial', size=10, bold=True)
    normal = Font(name='Arial', size=9)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    _merge(ws, 'A1:AF1', title, font=title_font, alignment=center)
    ws.row_dimensions[1].height = 34
    _merge(ws, 'A2:R2', 'Información General', fill=LIGHT_GRAY, font=section_font, alignment=center)
    _merge(ws, 'A3:A3', 'Empresa:', font=section_font, alignment=left)
    _merge(ws, 'B3:R3', meta.get('empresa',''), font=normal, alignment=left)
    _merge(ws, 'A4:A4', 'Área:', font=section_font, alignment=left)
    _merge(ws, 'B4:M4', meta.get('area',''), font=normal, alignment=left)
    _merge(ws, 'N4:O4', 'Fecha:', font=section_font, alignment=center)
    _merge(ws, 'P4:R4', meta.get('fecha',''), font=normal, alignment=center)
    _merge(ws, 'A5:A5', 'Subárea:', font=section_font, alignment=left)
    _merge(ws, 'B5:R5', meta.get('subarea',''), font=normal, alignment=left)
    _merge(ws, 'A6:R6', 'Condiciones de la actividad', fill=LIGHT_GRAY, font=section_font, alignment=center)
    _merge(ws, 'A7:A7', 'Puesto de trabajo:', font=section_font, alignment=left)
    _merge(ws, 'B7:R7', meta.get('puesto',''), font=normal, alignment=left)
    _merge(ws, 'A8:A8', 'Actividad:', font=section_font, alignment=left)
    _merge(ws, 'B8:R8', meta.get('actividad',''), font=normal, alignment=left)
    _merge(ws, 'A9:A9', 'Descripción de actividad:', font=section_font, alignment=left)
    _merge(ws, 'B9:R9', meta.get('descripcion',''), font=normal, alignment=left)
    ws.row_dimensions[9].height = 42

    fields = (payload or {}).get('snapshot',{}).get('fields',[]) or []
    def fv(*keys):
        keys = [str(k).lower() for k in keys]
        for f in fields:
            corpus = ' '.join(str(f.get(k,'')) for k in ('name','id','label')).lower()
            if any(k in corpus for k in keys):
                v = f.get('value')
                if v not in (None,''):
                    return v
        return ''

    _merge(ws, 'A10:D10', '# Trabajadores:', font=section_font, alignment=left)
    _merge(ws, 'E10:J10', fv('trabajadores','numero de trabajadores','# trabajadores'), font=normal, alignment=left)
    _merge(ws, 'K10:M10', 'Frecuencia/día:', font=section_font, alignment=left)
    _merge(ws, 'N10:R10', fv('frecuencia'), font=normal, alignment=left)
    _merge(ws, 'A11:D11', 'Tiempo de duración:', font=section_font, alignment=left)
    _merge(ws, 'E11:J11', fv('duración','duracion','tiempo'), font=normal, alignment=left)
    _merge(ws, 'K11:M11', 'Equipo auxiliar que utiliza:', font=section_font, alignment=left)
    _merge(ws, 'N11:R11', fv('equipo auxiliar','equipo'), font=normal, alignment=left)
    _merge(ws, 'A12:R12', 'Movimientos que realiza', fill=LIGHT_GRAY, font=section_font, alignment=center)
    movement_cols = [('A13:C13','Levantar','levantar'),('D13:F13','Bajar','bajar'),('G13:I13','Transportar','transportar'),('J13:L13','Empujar','empujar'),('M13:O13','Jalar','jalar'),('P13:R13','Estibar','estibar')]
    for rng,label,key in movement_cols:
        _merge(ws, rng, label, font=section_font, alignment=center)
        val = fv(key)
        checked = str(val).lower() in ('true','1','si','sí','on','x')
        start = rng.split(':')[0]
        col = start[0]
        # second line is merged same width; simple explicit mappings
    for rng in ['A14:C14','D14:F14','G14:I14','J14:L14','M14:O14','P14:R14']:
        _merge(ws, rng, '', font=normal, alignment=center)


def _risk_reference(ws, start_row=36):
    h = Font(name='Arial', size=9, bold=True)
    n = Font(name='Arial', size=8)
    c = Alignment(horizontal='center', vertical='center', wrap_text=True)
    l = Alignment(horizontal='left', vertical='center', wrap_text=True)
    _merge(ws, f'A{start_row}:J{start_row}', 'VALORES PARA LA DETERMINACIÓN DEL GRADO DE RIESGO', fill=LIGHT_GRAY, font=h, alignment=c)
    _merge(ws, f'A{start_row+1}:D{start_row+1}', 'NIVEL DE RIESGO', font=h, alignment=c)
    _merge(ws, f'E{start_row+1}:J{start_row+1}', 'PUNTAJE TOTAL', font=h, alignment=c)
    _merge(ws, f'K{start_row+1}:AF{start_row+1}', 'ACCIONES', font=h, alignment=c)
    rows = [
        ('Bajo – Aceptable','0 a 4','Sólo se requiere dar seguimiento a los grupos vulnerables como mujeres en periodo de gestación o trabajadores menores de edad.',GREEN,BLACK),
        ('Medio – Posible','5 a 12','Se debe examinar las tareas con mayor detalle, mediante una evaluación específica, o implantar medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.',YELLOW,BLACK),
        ('Alto – Significativo','13 a 20','Se requiere una acción rápida, por lo que se deben establecer medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.',RED,WHITE),
        ('Muy Alto – Inaceptable','21 a 32','Se deben detener las actividades e implementar medidas de control mediante un Programa de ergonomía para el manejo manual de cargas.',MAGENTA,WHITE),
    ]
    for i,(level,points,action,fill,font_color) in enumerate(rows,start_row+2):
        _merge(ws, f'A{i}:D{i}', level, fill=fill, font=Font(name='Arial',size=8,bold=True,color=font_color), alignment=c)
        _merge(ws, f'E{i}:J{i}', points, font=n, alignment=c)
        _merge(ws, f'K{i}:AF{i}', action, font=n, alignment=l)
        ws.row_dimensions[i].height = 34


def _add_justifications(ws, matrix, start_row):
    h = Font(name='Arial', size=10, bold=True)
    n = Font(name='Arial', size=8)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    rows = matrix.get('justification_rows') or []
    if not rows:
        return start_row
    _merge(ws, f'A{start_row}:AF{start_row}', 'FUNDAMENTO DE LA PUNTUACIÓN Y EVIDENCIA', fill=LIGHT_GRAY, font=h, alignment=center)
    start_row += 1
    _merge(ws, f'A{start_row}:H{start_row}', 'FACTOR / ACTIVIDAD', font=h, alignment=center)
    _merge(ws, f'I{start_row}:P{start_row}', 'CONDICIÓN SELECCIONADA', font=h, alignment=center)
    _merge(ws, f'Q{start_row}:R{start_row}', 'VALOR', font=h, alignment=center)
    _merge(ws, f'S{start_row}:Z{start_row}', 'JUSTIFICACIÓN DEL EVALUADOR', font=h, alignment=center)
    _merge(ws, f'AA{start_row}:AF{start_row}', 'EVIDENCIA', font=h, alignment=center)
    for item in rows:
        start_row += 1
        _merge(ws, f'A{start_row}:H{start_row}', item.get('title',''), font=n, alignment=left)
        _merge(ws, f'I{start_row}:P{start_row}', item.get('condition',''), font=n, alignment=left)
        _merge(ws, f'Q{start_row}:R{start_row}', item.get('value','—'), font=n, alignment=center)
        _merge(ws, f'S{start_row}:Z{start_row}', item.get('justification',''), font=n, alignment=left)
        photos = item.get('photos') or []
        names = ', '.join(p.get('name','Evidencia') for p in photos)
        _merge(ws, f'AA{start_row}:AF{start_row}', names or '—', font=n, alignment=left)
        ws.row_dimensions[start_row].height = 42
    return start_row


def build_appendix_i(matrix, payload):
    wb, ws = _base_sheet('APENDICE I')
    _general_header(ws, 'EVALUACIÓN DE RIESGO\n(LEVANTAMIENTO/DESCENSO, TRANSPORTE INDIVIDUAL Y EN EQUIPO)', matrix.get('meta',{}), payload)
    h = Font(name='Arial', size=9, bold=True)
    n = Font(name='Arial', size=8)
    c = Alignment(horizontal='center', vertical='center', wrap_text=True)
    l = Alignment(horizontal='left', vertical='center', wrap_text=True)
    _merge(ws, 'A15:AF15', 'EVALUACIÓN DE RIESGO POR MANEJO MANUAL DE CARGAS LEVANTAMIENTO/DESCENSO, TRANSPORTE INDIVIDUAL Y EN EQUIPO (APÉNDICE I NOM-036-1-STPS-2018)', font=h, alignment=c)
    _merge(ws, 'A16:M17', 'FACTORES DE RIESGO', fill=LIGHT_GRAY, font=h, alignment=c)
    _merge(ws, 'N16:V16', 'LEVANTAR', fill=LIGHT_GRAY, font=h, alignment=c)
    _merge(ws, 'W16:AB16', 'TRANSPORTAR', fill=LIGHT_GRAY, font=h, alignment=c)
    _merge(ws, 'AC16:AF16', 'EQUIPO', fill=LIGHT_GRAY, font=h, alignment=c)
    for rng,val in [('N17:Q17','COLOR'),('R17:V17','VALOR'),('W17:Y17','COLOR'),('Z17:AB17','VALOR'),('AC17:AE17','COLOR'),('AF17:AF17','VALOR')]:
        _merge(ws,rng,val,fill=LIGHT_GRAY,font=h,alignment=c)

    matrix_rows = matrix.get('factor_rows') or []
    for idx,item in enumerate(matrix_rows, start=18):
        _merge(ws, f'A{idx}:M{idx}', item.get('label',''), font=n, alignment=l)
        cells = item.get('cells') or [{},{},{}]
        blocks = [('N','R'),('W','Z'),('AC','AF')]
        color_ranges = [('N','Q'),('W','Y'),('AC','AE')]
        value_ranges = [('R','V'),('Z','AB'),('AF','AF')]
        for j,cell in enumerate(cells[:3]):
            cr = f'{color_ranges[j][0]}{idx}:{color_ranges[j][1]}{idx}'
            vr = f'{value_ranges[j][0]}{idx}:{value_ranges[j][1]}{idx}'
            _merge(ws, cr, cell.get('color') or '—', font=n, alignment=c)
            top = ws[cr.split(':')[0]]
            color_name = cell.get('color') or ''
            if color_name in RISK_STYLES:
                fill,font_color = RISK_STYLES[color_name]
                top.fill = PatternFill('solid',fgColor=fill)
                top.font = Font(name='Arial',size=8,bold=True,color=font_color)
            _merge(ws, vr, cell.get('value') if cell.get('value') is not None else '—', font=n, alignment=c)
        ws.row_dimensions[idx].height = 26

    scores = matrix.get('scores') or {}
    levels = matrix.get('levels') or {}
    for labelrow, label in [(29,'PUNTUACIÓN'),(30,'NIVEL DE RIESGO'),(31,'ACCIONES')]:
        _merge(ws, f'A{labelrow}:M{labelrow}', label, font=h, alignment=c)
    for sec,cr in [('ai2','N:V'),('ai3','W:AB'),('ai4','AC:AF')]:
        a,b = cr.split(':')
        _merge(ws, f'{a}29:{b}29', scores.get(sec,0), font=h, alignment=c)
        lev, action = _risk_label(scores.get(sec,0))
        _merge(ws, f'{a}30:{b}30', lev, font=h, alignment=c)
        _merge(ws, f'{a}31:{b}32', action, font=n, alignment=l)
    ws.row_dimensions[31].height = 34
    ws.row_dimensions[32].height = 34

    _risk_reference(ws, 36)
    last = _add_justifications(ws, matrix, 42)
    ws.print_area = f'A1:AF{max(last,41)}'
    return wb


def _snapshot_value(payload, *needles):
    needles = [n.lower() for n in needles]
    fields = (payload.get('snapshot') or {}).get('fields') or []
    for f in fields:
        corpus = ' '.join(str(f.get(k,'')) for k in ('name','id','label')).lower()
        if any(n in corpus for n in needles):
            return f.get('value')
    return None


def build_appendix_ii(matrix, payload):
    mode = str(_snapshot_value(payload,'modo','modalidad') or '').lower()
    con = mode in ('con','con equipo','con equipo auxiliar') or 'con equipo' in mode
    sheet_title = 'APENDICE II'
    wb, ws = _base_sheet(sheet_title)
    title = 'EVALUACIÓN DE RIESGO\n(EMPUJAR O JALAR CARGAS CON USO DE EQUIPO AUXILIAR)' if con else 'EVALUACIÓN DE RIESGO\n(EMPUJAR O JALAR CARGAS SIN USO DE EQUIPO AUXILIAR)'
    _general_header(ws, title, matrix.get('meta',{}), payload)
    h = Font(name='Arial', size=9, bold=True)
    n = Font(name='Arial', size=8)
    c = Alignment(horizontal='center', vertical='center', wrap_text=True)
    l = Alignment(horizontal='left', vertical='center', wrap_text=True)
    _merge(ws, 'A15:AF15', 'EVALUACIÓN DE RIESGO POR EMPUJE Y ARRASTRE DE CARGAS ' + ('CON EQUIPO AUXILIAR' if con else 'SIN EQUIPO AUXILIAR') + ' (APÉNDICE II NOM-036-1-STPS-2018)', font=h, alignment=c)
    _merge(ws, 'A16:M17', 'FACTORES DE RIESGO', fill=LIGHT_GRAY, font=h, alignment=c)
    if con:
        headers = ['EQUIPO PEQUEÑO','EQUIPO MEDIANO','EQUIPO GRANDE']
        selected = str(_snapshot_value(payload,'equipo') or '').lower()
        sel_idx = 0 if 'peq' in selected or 'peque' in selected else (1 if 'med' in selected else (2 if 'grand' in selected else 0))
    else:
        headers = ['RODANDO','GIRANDO SOBRE SU BASE','ARRASTRANDO/JALANDO O DESLIZANDO']
        selected = str(_snapshot_value(payload,'mov','desplazamiento') or '').lower()
        sel_idx = 0 if 'rod' in selected else (1 if 'gir' in selected else (2 if 'arras' in selected or 'jal' in selected or 'desliz' in selected else 0))
    _merge(ws,'N16:V16',headers[0],fill=LIGHT_GRAY,font=h,alignment=c)
    _merge(ws,'W16:AB16',headers[1],fill=LIGHT_GRAY,font=h,alignment=c)
    _merge(ws,'AC16:AF16',headers[2],fill=LIGHT_GRAY,font=h,alignment=c)
    for rng,val in [('N17:Q17','NIVEL DE RIESGO'),('R17:V17','VALOR'),('W17:Y17','NIVEL DE RIESGO'),('Z17:AB17','VALOR'),('AC17:AE17','NIVEL DE RIESGO'),('AF17:AF17','VALOR')]:
        _merge(ws,rng,val,fill=LIGHT_GRAY,font=h,alignment=c)

    factors = matrix.get('factors') or []
    row = 18
    for item in factors:
        # Ap II without equipment omits equipment-condition row.
        if not con and 'equipo auxiliar' in item.get('title','').lower():
            continue
        _merge(ws,f'A{row}:M{row}',item.get('title',''),font=n,alignment=l)
        value = item.get('value')
        color = item.get('color') or ''
        blocks=[('N','Q','R','V'),('W','Y','Z','AB'),('AC','AE','AF','AF')]
        for idx,(c1,c2,v1,v2) in enumerate(blocks):
            if idx == sel_idx:
                level_text = item.get('condition') or color or ('BAJO' if value in (0,None) else '')
                _merge(ws,f'{c1}{row}:{c2}{row}',level_text,font=n,alignment=c)
                top=ws[f'{c1}{row}']
                if color in RISK_STYLES:
                    fill,font_color=RISK_STYLES[color]
                    top.fill=PatternFill('solid',fgColor=fill)
                    top.font=Font(name='Arial',size=8,bold=True,color=font_color)
                _merge(ws,f'{v1}{row}:{v2}{row}',value if value is not None else '—',font=n,alignment=c)
            else:
                _merge(ws,f'{c1}{row}:{c2}{row}','NO APLICA',fill=GRAY,font=Font(name='Arial',size=8,bold=True,color=WHITE),alignment=c)
                _merge(ws,f'{v1}{row}:{v2}{row}','—',fill=GRAY,font=Font(name='Arial',size=8,bold=True,color=WHITE),alignment=c)
        row += 1
    total_row = row
    level_row = row+1
    action_row = row+2
    _merge(ws,f'A{total_row}:M{total_row}','PUNTUACIÓN',font=h,alignment=c)
    _merge(ws,f'A{level_row}:M{level_row}','NIVEL DE RIESGO',font=h,alignment=c)
    _merge(ws,f'A{action_row}:M{action_row}','ACCIONES',font=h,alignment=c)
    score = matrix.get('final_score') or 0
    level,action = _risk_label(score)
    blocks=[('N','V'),('W','AB'),('AC','AF')]
    for idx,(a,b) in enumerate(blocks):
        if idx==sel_idx:
            _merge(ws,f'{a}{total_row}:{b}{total_row}',score,font=h,alignment=c)
            _merge(ws,f'{a}{level_row}:{b}{level_row}',level,font=h,alignment=c)
            _merge(ws,f'{a}{action_row}:{b}{action_row+1}',action,font=n,alignment=l)
        else:
            _merge(ws,f'{a}{total_row}:{b}{total_row}','—',fill=GRAY,font=Font(name='Arial',size=8,bold=True,color=WHITE),alignment=c)
            _merge(ws,f'{a}{level_row}:{b}{level_row}','NO APLICA',fill=GRAY,font=Font(name='Arial',size=8,bold=True,color=WHITE),alignment=c)
            _merge(ws,f'{a}{action_row}:{b}{action_row+1}','NO APLICA',fill=GRAY,font=Font(name='Arial',size=8,bold=True,color=WHITE),alignment=c)
    ref_start=action_row+5
    _risk_reference(ws, ref_start)
    last=_add_justifications(ws,matrix,ref_start+6)
    ws.print_area=f'A1:AF{max(last,ref_start+5)}'
    return wb


def export_nom_excel(matrix, payload):
    if matrix.get('kind') == 'APENDICE_I':
        wb = build_appendix_i(matrix,payload)
        filename = f"Fo-NSTPS-61_{matrix.get('folio','APENDICE_I')}.xlsx"
    elif matrix.get('kind') == 'APENDICE_II':
        wb = build_appendix_ii(matrix,payload)
        mode = str(_snapshot_value(payload,'modo','modalidad') or '').lower()
        con = mode in ('con','con equipo','con equipo auxiliar') or 'con equipo' in mode
        code = 'Fo-NSTPS-63' if con else 'Fo-NSTPS-62'
        filename = f"{code}_{matrix.get('folio','APENDICE_II')}.xlsx"
    else:
        raise ValueError('Método sin plantilla Excel configurada')
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, filename
